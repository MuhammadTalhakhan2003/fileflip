"""Convert from PDF: Word, JPG/PNG, Excel, PowerPoint, text, Markdown, PDF/A."""

import logging
import re
import subprocess
import threading

import pymupdf
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter

from app.core import engines
from app.core.errors import ConversionFailed, InvalidInput
from app.tools.common import PAGES, PASSWORD, bundle, open_pdf, out_path, parse_page_ranges
from app.tools.registry import Category, Context, Option, Result, tool

log = logging.getLogger(__name__)
logging.getLogger("pdf2docx").setLevel(logging.ERROR)


# ---------------------------------------------------------------- Word


@tool(
    slug="pdf-to-word",
    name="PDF to Word",
    category=Category.FROM_PDF,
    description="Turn PDFs into editable Word documents, keeping the layout.",
    accepts=("pdf",),
    output="docx",
    options=(PAGES, PASSWORD),
    aliases=("pdf-to-docx", "pdf-to-doc"),
    timeout=600,
)
def pdf_to_word(ctx: Context) -> Result:
    from pdf2docx import Converter

    src = ctx.inputs[0]
    password = ctx.options["password"] or None
    with open_pdf(src, ctx.options["password"]) as doc:
        pages = parse_page_ranges(ctx.options["pages"], doc.page_count)
        has_text = any(doc[p].get_text("text").strip() for p in pages[:20])

    out = out_path(ctx.work_dir, src.stem, ".docx")
    ctx.progress(0.04, "Opening document")
    listener = _Pdf2DocxProgress(ctx.progress)
    root = logging.getLogger()
    root.addHandler(listener)
    converter = Converter(str(src), password=password)
    try:
        converter.convert(str(out), pages=pages)
    except Exception as exc:
        raise ConversionFailed("We couldn't convert this PDF to Word. The layout may be too complex.") from exc
    finally:
        root.removeHandler(listener)
        converter.close()

    meta = {"pages": len(pages)}
    if not has_text:
        meta["warning"] = "This looks like a scanned PDF. Run 'OCR PDF' first to get editable text."
    return Result(out, meta)


class _Pdf2DocxProgress(logging.Handler):
    """pdf2docx logs '[3/4] Parsing pages...' then '(2/9) Page 2' for every page - turn that into progress."""

    STAGES = {
        "1": (0.05, 0.08, "Opening document"),
        "2": (0.08, 0.12, "Analyzing layout"),
        "3": (0.12, 0.60, "Reading page {n} of {total}"),
        "4": (0.60, 0.95, "Building Word page {n} of {total}"),
    }

    def __init__(self, report):
        super().__init__(logging.INFO)
        self.report = report
        self.thread = threading.get_ident()  # other conversions may be logging at the same time
        self.stage = self.STAGES["1"]

    def emit(self, record: logging.LogRecord) -> None:
        if record.thread != self.thread:
            return
        msg = record.getMessage()
        if m := re.search(r"\[([1-4])/4\]", msg):
            self.stage = self.STAGES[m.group(1)]
            start, _, label = self.stage
            self.report(start, label.split(" {n}")[0].replace("page", "pages") if "{n}" in label else label)
        elif m := re.match(r"\((\d+)/(\d+)\) Page", msg):
            n, total = int(m.group(1)), int(m.group(2))
            start, end, label = self.stage
            self.report(start + (end - start) * n / max(total, 1), label.format(n=n, total=total))


# ---------------------------------------------------------------- images


def _pdf_to_images(ctx: Context, fmt: str) -> Result:
    src = ctx.inputs[0]
    dpi = ctx.options["dpi"]
    outputs = []
    with open_pdf(src, ctx.options["password"]) as doc:
        pages = parse_page_ranges(ctx.options["pages"], doc.page_count)
        parts = ctx.work_dir / "pages"
        parts.mkdir()
        for i in ctx.steps(len(pages), "Rendering page {n} of {total}"):
            p = pages[i]
            pix = doc[p].get_pixmap(dpi=dpi, alpha=False)
            dest = parts / f"{src.stem}_page_{p + 1}.{fmt}"
            if fmt == "jpg":
                pix.save(dest, jpg_quality=ctx.options["quality"])
            else:
                pix.save(dest)
            outputs.append(dest)
    return Result(bundle(outputs, ctx.work_dir, f"{src.stem}_{fmt}"), {"pages": len(outputs), "dpi": dpi})


IMAGE_OPTIONS = (
    Option("dpi", "int", default=150, label="Resolution (DPI)", min=50, max=300),
    PAGES,
    PASSWORD,
)


@tool(
    slug="pdf-to-jpg",
    name="PDF to JPG",
    category=Category.FROM_PDF,
    description="Convert each PDF page into a high-quality JPG image.",
    accepts=("pdf",),
    output="jpg",
    options=IMAGE_OPTIONS + (Option("quality", "int", default=90, label="JPG quality", min=30, max=100),),
    aliases=("pdf-to-jpeg", "pdf-to-image"),
)
def pdf_to_jpg(ctx: Context) -> Result:
    return _pdf_to_images(ctx, "jpg")


@tool(
    slug="pdf-to-png",
    name="PDF to PNG",
    category=Category.FROM_PDF,
    description="Convert each PDF page into a lossless PNG image.",
    accepts=("pdf",),
    output="png",
    options=IMAGE_OPTIONS,
)
def pdf_to_png(ctx: Context) -> Result:
    return _pdf_to_images(ctx, "png")


# ---------------------------------------------------------------- Excel

_NUMBER = re.compile(r"^-?\(?[$€£₹]?\s?\d{1,3}(,\d{3})*(\.\d+)?\)?%?$|^-?\d+(\.\d+)?%?$")


def _cell_value(raw):
    if raw is None:
        return None
    text = ILLEGAL_CHARACTERS_RE.sub("", str(raw)).strip()
    if not text:
        return None
    if _NUMBER.match(text):
        negative = text.startswith("(") and text.endswith(")") or text.startswith("-")
        digits = re.sub(r"[^\d.]", "", text)
        try:
            value = float(digits)
        except ValueError:
            return text
        if text.endswith("%"):
            value /= 100
        value = -value if negative else value
        return int(value) if value.is_integer() and "." not in digits and not text.endswith("%") else value
    return text


def _text_rows(page: pymupdf.Page) -> list[list[str]]:
    """Fallback when no ruled tables exist: rebuild rows from word positions,
    starting a new cell wherever there's a wide horizontal gap."""
    rows: dict[tuple[int, int], list] = {}
    for x0, y0, x1, y1, word, block, line, _ in page.get_text("words", sort=True):
        rows.setdefault((block, line), []).append((x0, x1, word, y1 - y0))
    result = []
    for words in rows.values():
        cells, current, last_x1 = [], [], None
        for x0, x1, word, h in words:
            if last_x1 is not None and x0 - last_x1 > max(h, 6) * 1.2:
                cells.append(" ".join(current))
                current = []
            current.append(word)
            last_x1 = x1
        cells.append(" ".join(current))
        result.append(cells)
    return result


def _table_rows(table, words: list) -> list[list[str | None]]:
    """Use the detected cell grid, but fill each cell from word positions - the
    built-in cell text extraction can scramble characters like decimal points."""
    rows = []
    for row in table.rows:
        values = []
        for cell in row.cells:
            if cell is None:
                values.append(None)
                continue
            x0, y0, x1, y1 = cell
            inside = [w for w in words if x0 <= (w[0] + w[2]) / 2 <= x1 and y0 <= (w[1] + w[3]) / 2 <= y1]
            lines: dict[tuple[int, int], list[str]] = {}
            for w in sorted(inside, key=lambda w: (w[5], w[6], w[7])):
                lines.setdefault((w[5], w[6]), []).append(w[4])
            values.append("\n".join(" ".join(parts) for parts in lines.values()) or None)
        rows.append(values)
    return rows


def _write_rows(ws, rows) -> None:
    widths: dict[int, int] = {}
    for row in rows:
        values = [_cell_value(c) for c in row]
        ws.append(values)
        for i, v in enumerate(values, start=1):
            widths[i] = max(widths.get(i, 0), min(60, len(str(v)) if v is not None else 0))
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = max(8, w + 2)


@tool(
    slug="pdf-to-excel",
    name="PDF to Excel",
    category=Category.FROM_PDF,
    description="Pull tables and data from PDFs straight into Excel spreadsheets.",
    accepts=("pdf",),
    output="xlsx",
    options=(
        Option("layout", "choice", default="sheet-per-table", label="Layout", choices=("sheet-per-table", "sheet-per-page", "single-sheet")),
        PAGES,
        PASSWORD,
    ),
    aliases=("pdf-to-xlsx",),
    timeout=600,
)
def pdf_to_excel(ctx: Context) -> Result:
    src = ctx.inputs[0]
    layout = ctx.options["layout"]
    wb = Workbook()
    wb.remove(wb.active)
    tables_found = 0

    def sheet(title: str):
        title = re.sub(r"[\[\]:*?/\\]", "", title)[:31]
        base, n = title, 2
        while title in wb.sheetnames:
            title = f"{base[:27]} ({n})"
            n += 1
        return wb.create_sheet(title)

    single = sheet("Data") if layout == "single-sheet" else None
    with open_pdf(src, ctx.options["password"]) as doc:
        pages = parse_page_ranges(ctx.options["pages"], doc.page_count)
        for i in ctx.steps(len(pages), "Finding tables on page {n} of {total}", end=0.9):
            p = pages[i]
            page = doc[p]
            try:
                tables = page.find_tables().tables
            except Exception:
                tables = []
            page_sheet = None
            words = page.get_text("words") if tables else []
            if not tables:
                rows = _text_rows(page)
                if rows:
                    target = single or sheet(f"Page {p + 1}")
                    _write_rows(target, rows)
                    if single:
                        single.append([])
                continue
            for t_index, table in enumerate(tables, start=1):
                tables_found += 1
                rows = _table_rows(table, words)
                if layout == "sheet-per-table":
                    target = sheet(f"Page {p + 1}" + (f" Table {t_index}" if len(tables) > 1 else ""))
                elif layout == "sheet-per-page":
                    page_sheet = page_sheet or sheet(f"Page {p + 1}")
                    target = page_sheet
                else:
                    target = single
                _write_rows(target, rows)
                if target.max_row and layout != "sheet-per-table":
                    target.append([])

    if not wb.sheetnames:
        raise InvalidInput("No text or tables were found. If this is a scanned PDF, run 'OCR PDF' first.")
    out = out_path(ctx.work_dir, src.stem, ".xlsx")
    wb.save(out)
    return Result(out, {"tables": tables_found, "sheets": len(wb.sheetnames)})


# ---------------------------------------------------------------- PowerPoint

_MAX_SLIDE_EMU = 51_206_400  # PowerPoint's 56 inch limit
_EMU_PER_PT = 12_700


@tool(
    slug="pdf-to-powerpoint",
    name="PDF to PowerPoint",
    category=Category.FROM_PDF,
    description="Turn PDFs into PowerPoint slides - one slide per page, with page text in the notes.",
    accepts=("pdf",),
    output="pptx",
    options=(Option("dpi", "int", default=200, label="Slide image quality (DPI)", min=96, max=300), PAGES, PASSWORD),
    aliases=("pdf-to-pptx", "pdf-to-ppt"),
)
def pdf_to_powerpoint(ctx: Context) -> Result:
    from pptx import Presentation

    src = ctx.inputs[0]
    prs = Presentation()
    blank = prs.slide_layouts[6]
    with open_pdf(src, ctx.options["password"]) as doc:
        pages = parse_page_ranges(ctx.options["pages"], doc.page_count)
        first = doc[pages[0]].rect
        scale = min(1.0, _MAX_SLIDE_EMU / (max(first.width, first.height) * _EMU_PER_PT))
        prs.slide_width = int(first.width * _EMU_PER_PT * scale)
        prs.slide_height = int(first.height * _EMU_PER_PT * scale)

        img_dir = ctx.work_dir / "slides"
        img_dir.mkdir()
        for i in ctx.steps(len(pages), "Making slide {n} of {total}", end=0.9):
            p = pages[i]
            page = doc[p]
            img = img_dir / f"{p}.png"
            page.get_pixmap(dpi=ctx.options["dpi"], alpha=False).save(img)
            slide = prs.slides.add_slide(blank)
            # fit each page inside the slide, centered (pages may differ in size)
            fit = min(prs.slide_width / (page.rect.width * _EMU_PER_PT), prs.slide_height / (page.rect.height * _EMU_PER_PT))
            w = int(page.rect.width * _EMU_PER_PT * fit)
            h = int(page.rect.height * _EMU_PER_PT * fit)
            slide.shapes.add_picture(str(img), (prs.slide_width - w) // 2, (prs.slide_height - h) // 2, w, h)
            text = ILLEGAL_CHARACTERS_RE.sub("", page.get_text("text", sort=True)).strip()
            if text:
                slide.notes_slide.notes_text_frame.text = text[:20_000]

    out = out_path(ctx.work_dir, src.stem, ".pptx")
    prs.save(out)
    return Result(out, {"slides": len(pages)})


# ---------------------------------------------------------------- text / markdown


@tool(
    slug="pdf-to-text",
    name="PDF to Text",
    category=Category.FROM_PDF,
    description="Extract all the text from a PDF into a plain .txt file.",
    accepts=("pdf",),
    output="txt",
    options=(PAGES, PASSWORD),
    aliases=("pdf-to-txt",),
)
def pdf_to_text(ctx: Context) -> Result:
    src = ctx.inputs[0]
    with open_pdf(src, ctx.options["password"]) as doc:
        pages = parse_page_ranges(ctx.options["pages"], doc.page_count)
        chunks = [doc[p].get_text("text", sort=True).strip() for p in pages]
    text = "\n\n\f".join(chunks).strip()
    if not text:
        raise InvalidInput("No text was found. If this is a scanned PDF, run 'OCR PDF' first.")
    out = out_path(ctx.work_dir, src.stem, ".txt")
    out.write_text(text, encoding="utf-8")
    return Result(out, {"characters": len(text), "pages": len(pages)})


@tool(
    slug="pdf-to-markdown",
    name="PDF to Markdown",
    category=Category.INTELLIGENCE,
    description="Convert PDFs to clean Markdown with headings, lists and tables - ready for AI tools and docs.",
    accepts=("pdf",),
    output="md",
    options=(PAGES, PASSWORD),
    aliases=("pdf-to-md",),
    timeout=600,
)
def pdf_to_markdown(ctx: Context) -> Result:
    src = ctx.inputs[0]
    with open_pdf(src, ctx.options["password"]) as doc:
        pages = parse_page_ranges(ctx.options["pages"], doc.page_count)
        try:
            import pymupdf4llm

            # a few pages per call, so progress can move
            batches = [pages[i:i + 5] for i in range(0, len(pages), 5)]
            md = "".join(
                pymupdf4llm.to_markdown(doc, pages=batches[i], show_progress=False)
                for i in ctx.steps(len(batches), "Reading part {n} of {total}")
            )
        except Exception:
            log.warning("pymupdf4llm failed, falling back to plain text", exc_info=True)
            md = "\n\n---\n\n".join(doc[p].get_text("text", sort=True) for p in pages)
    if not md.strip():
        raise InvalidInput("No text was found. If this is a scanned PDF, run 'OCR PDF' first.")
    out = out_path(ctx.work_dir, src.stem, ".md")
    out.write_text(md, encoding="utf-8")
    return Result(out, {"characters": len(md), "pages": len(pages)})


# ---------------------------------------------------------------- PDF/A


_PDFA_DEF = """%!
/ICCProfile ({icc}) def
[/_objdef {{icc_PDFA}} /type /stream /OBJ pdfmark
[{{icc_PDFA}} <</N 3>> /PUT pdfmark
[{{icc_PDFA}} ICCProfile (r) file /PUT pdfmark
[/_objdef {{OutputIntent_PDFA}} /type /dict /OBJ pdfmark
[{{OutputIntent_PDFA}} <<
  /Type /OutputIntent
  /S /GTS_PDFA1
  /DestOutputProfile {{icc_PDFA}}
  /OutputConditionIdentifier (sRGB)
>> /PUT pdfmark
[{{Catalog}} <</OutputIntents [ {{OutputIntent_PDFA}} ]>> /PUT pdfmark
"""


@tool(
    slug="pdf-to-pdfa",
    name="PDF to PDF/A",
    category=Category.FROM_PDF,
    description="Convert PDFs to PDF/A, the ISO standard for long-term archiving.",
    accepts=("pdf",),
    output="pdf",
    requires=(engines.GHOSTSCRIPT,),
    options=(Option("conformance", "choice", default="2b", label="Conformance", choices=("1b", "2b", "3b")), PASSWORD),
    timeout=600,
)
def pdf_to_pdfa(ctx: Context) -> Result:
    from PIL import ImageCms

    src = ctx.inputs[0]
    # Ghostscript can't open encrypted files, so decrypt first if needed.
    with open_pdf(src, ctx.options["password"]) as doc:
        if doc.is_encrypted or ctx.options["password"]:
            plain = ctx.work_dir / "decrypted.pdf"
            doc.save(plain, encryption=pymupdf.PDF_ENCRYPT_NONE)
            src_for_gs = plain
        else:
            src_for_gs = src

    icc = ctx.work_dir / "srgb.icc"
    icc.write_bytes(ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes())
    icc_ps = icc.as_posix().replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    pdfa_def = ctx.work_dir / "PDFA_def.ps"
    pdfa_def.write_text(_PDFA_DEF.format(icc=icc_ps), encoding="latin-1")

    out = out_path(ctx.work_dir, f"{src.stem}_pdfa", ".pdf")
    level = ctx.options["conformance"][0]
    cmd = [
        engines.ghostscript_path(),
        f"-dPDFA={level}",
        "-dBATCH",
        "-dNOPAUSE",
        "-dNOOUTERSAVE",
        "-dSAFER",
        f"--permit-file-read={icc.as_posix()}",
        "-sColorConversionStrategy=RGB",
        "-dPDFACompatibilityPolicy=1",
        "-sDEVICE=pdfwrite",
        f"-sOutputFile={out}",
        str(pdfa_def),
        str(src_for_gs),
    ]
    proc = subprocess.run(cmd, capture_output=True, timeout=540)
    if proc.returncode != 0 or not out.is_file():
        log.warning("Ghostscript PDF/A failed: %s", proc.stderr.decode(errors="ignore")[-2000:])
        raise ConversionFailed("We couldn't convert this PDF to PDF/A.")
    return Result(out, {"conformance": f"PDF/A-{ctx.options['conformance'].upper()}"})
