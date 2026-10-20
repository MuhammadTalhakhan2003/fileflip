import io
import json
import zipfile

import pymupdf
import pytest
from docx import Document as DocxDocument
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation

from app.core import engines
from tests.conftest import make_form_pdf, make_image, make_pdf

PDF = "application/pdf"


def run(client, slug, files, data=None, extras=None, expect=200):
    multipart = [("files", f) for f in files]
    for field, f in (extras or {}).items():
        multipart.append((field, f))
    resp = client.post(f"/api/v1/convert/{slug}", files=multipart, data=data or {})
    assert resp.status_code == expect, resp.text
    if expect != 200:
        return resp.json()
    record = resp.json()
    dl = client.get(record["download_url"])
    assert dl.status_code == 200
    assert "attachment" in dl.headers["content-disposition"]
    return record, dl.content


def pdf_file(name="doc.pdf", **kw):
    return (name, make_pdf(**kw), PDF)


def open_bytes(data) -> pymupdf.Document:
    return pymupdf.open(stream=data, filetype="pdf")


# ---------------------------------------------------------------- catalog / health


def test_tools_catalog(client):
    body = client.get("/api/v1/tools").json()
    slugs = {t["slug"] for c in body["categories"] for t in c["tools"]}
    for slug in ("merge-pdf", "pdf-to-word", "jpg-to-png", "sign-pdf", "translate-pdf", "pdf-to-pdfa"):
        assert slug in slugs
    assert client.get("/api/v1/tools/pdf-to-docx").json()["slug"] == "pdf-to-word"  # alias
    assert client.get("/api/v1/health").json()["status"] == "ok"


# ---------------------------------------------------------------- organize


def test_merge(client):
    record, data = run(client, "merge-pdf", [pdf_file("a.pdf", pages=2), pdf_file("b.pdf", pages=3)])
    assert open_bytes(data).page_count == 5
    assert record["meta"]["pages"] == 5


def test_split_ranges_and_all(client):
    _, data = run(client, "split-pdf", [pdf_file(pages=6)], {"mode": "ranges", "ranges": "1-2,3-6"})
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    assert len(names) == 2
    _, data = run(client, "split-pdf", [pdf_file(pages=4)], {"mode": "every", "every": "3"})
    assert len(zipfile.ZipFile(io.BytesIO(data)).namelist()) == 2


def test_remove_and_extract(client):
    _, data = run(client, "remove-pages", [pdf_file(pages=5)], {"pages": "2,4"})
    assert open_bytes(data).page_count == 3
    _, data = run(client, "extract-pages", [pdf_file(pages=5)], {"pages": "1,5"})
    doc = open_bytes(data)
    assert doc.page_count == 2 and "Hello page 5" in doc[1].get_text()
    err = run(client, "remove-pages", [pdf_file(pages=2)], {"pages": "1-2"}, expect=422)
    assert err["error"]["code"] == "invalid_input"


def test_organize(client):
    order = [{"file": 0, "page": 3, "rotate": 90}, {"blank": True}, {"file": 1, "page": 1}]
    _, data = run(client, "organize-pdf", [pdf_file("a.pdf"), pdf_file("b.pdf", text="Other")],
                  {"order": json.dumps(order)})
    doc = open_bytes(data)
    assert doc.page_count == 3
    assert doc[0].rotation == 90 and "Hello page 3" in doc[0].get_text()
    assert doc[1].get_text().strip() == ""
    assert "Other 1" in doc[2].get_text()


# ---------------------------------------------------------------- optimize


def test_compress(client):
    big = pymupdf.open()
    page = big.new_page()
    img = Image.effect_noise((2400, 2400), 60).convert("RGB")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    page.insert_image(page.rect, stream=buf.getvalue())
    record, data = run(client, "compress-pdf", [("big.pdf", big.tobytes(), PDF)], {"level": "extreme"})
    assert record["meta"]["compressed_size"] < record["meta"]["original_size"]
    assert open_bytes(data).page_count == 1


def test_repair(client):
    raw = make_pdf(pages=2)
    damaged = raw.replace(b"xref", b"xxxx", 1)  # break the cross-reference table
    _, data = run(client, "repair-pdf", [("broken.pdf", damaged, PDF)])
    assert open_bytes(data).page_count == 2


# ---------------------------------------------------------------- from PDF


def test_pdf_to_word(client):
    _, data = run(client, "pdf-to-word", [pdf_file(pages=2)])
    text = "\n".join(p.text for p in DocxDocument(io.BytesIO(data)).paragraphs)
    assert "Hello page 1" in text


def test_pdf_to_jpg_and_png(client):
    _, data = run(client, "pdf-to-jpg", [pdf_file(pages=1)], {"dpi": "72"})
    assert Image.open(io.BytesIO(data)).format == "JPEG"
    _, data = run(client, "pdf-to-png", [pdf_file(pages=3)], {"pages": "1-2"})
    assert len(zipfile.ZipFile(io.BytesIO(data)).namelist()) == 2


def test_pdf_to_excel(client):
    record, data = run(client, "pdf-to-excel", [pdf_file(pages=1, table=True)])
    wb = load_workbook(io.BytesIO(data))
    rows = [[c.value for c in row] for row in wb.worksheets[0].iter_rows()]
    assert rows[0] == ["Item", "Qty", "Price"]
    assert rows[1] == ["Row 1", 2, 10.5]  # numbers become real numbers, not text
    assert record["meta"]["tables"] >= 1


def test_pdf_to_powerpoint(client):
    _, data = run(client, "pdf-to-powerpoint", [pdf_file(pages=2)], {"dpi": "96"})
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 2
    assert "Hello page 1" in prs.slides[0].notes_slide.notes_text_frame.text


def test_pdf_to_text_and_markdown(client):
    _, data = run(client, "pdf-to-text", [pdf_file(pages=2)])
    assert "Hello page 2" in data.decode()
    _, data = run(client, "pdf-to-markdown", [pdf_file(pages=1)])
    assert "Hello page 1" in data.decode()


# ---------------------------------------------------------------- to PDF / images


def test_jpg_to_pdf(client):
    files = [("a.jpg", make_image("JPEG", mode="RGB"), "image/jpeg"), ("b.png", make_image("PNG"), "image/png")]
    _, data = run(client, "jpg-to-pdf", files, {"page_size": "a4", "margin": "small"})
    doc = open_bytes(data)
    assert doc.page_count == 2
    assert round(doc[0].rect.width) == 842  # landscape image -> landscape A4 in auto mode


def test_scan_to_pdf(client):
    _, data = run(client, "scan-to-pdf", [("scan.jpg", make_image("JPEG", mode="RGB"), "image/jpeg")],
                  {"color": "bw"})
    assert open_bytes(data).page_count == 1


def test_html_to_pdf(client):
    html = b"<html><body><h1>Invoice</h1><p>Total: 42</p></body></html>"
    _, data = run(client, "html-to-pdf", [("page.html", html, "text/html")])
    assert "Invoice" in open_bytes(data)[0].get_text()


@pytest.mark.parametrize("slug,src,fmt", [
    ("jpg-to-png", ("a.jpg", "JPEG", "RGB"), "PNG"),
    ("png-to-jpg", ("a.png", "PNG", "RGBA"), "JPEG"),
    ("png-to-webp", ("a.png", "PNG", "RGBA"), "WEBP"),
    ("webp-to-jpg", ("a.webp", "WEBP", "RGB"), "JPEG"),
    ("png-to-ico", ("a.png", "PNG", "RGBA"), "ICO"),
])
def test_image_pairs(client, slug, src, fmt):
    name, pil_fmt, mode = src
    _, data = run(client, slug, [(name, make_image(pil_fmt, mode=mode), "image/x")])
    assert Image.open(io.BytesIO(data)).format == fmt


SVG = b"""<svg xmlns="http://www.w3.org/2000/svg" width="120" height="60" viewBox="0 0 120 60">
<rect width="120" height="60" rx="10" fill="#5145f0"/><circle cx="30" cy="30" r="18" fill="#ffffff"/></svg>"""

XPM = b"""/* XPM */
static char *icon[] = {
"4 2 2 1",
"  c #FFFFFF",
". c #FF0000",
" .. ",
"....",
};
"""


def test_svg_to_png_is_rasterized_at_2x(client):
    _, data = run(client, "svg-to-png", [("logo.svg", SVG, "image/svg+xml")])
    img = Image.open(io.BytesIO(data))
    assert img.format == "PNG" and img.size == (240, 120)
    assert img.convert("RGB").getpixel((5, 60))[2] > 200  # brand blue survived


def test_unix_image_formats(client):
    def encode(fmt, mode):
        buf = io.BytesIO()
        Image.new(mode, (16, 8), 1 if mode == "1" else (200, 30, 30)).save(buf, fmt)
        return buf.getvalue()

    files = [("a.ppm", encode("PPM", "RGB"), "image/x"), ("b.xbm", encode("XBM", "1"), "image/x"), ("c.xpm", XPM, "image/x")]
    _, data = run(client, "image-converter", files, {"format": "png"})
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    assert sorted(names) == ["a.png", "b.png", "c.png"]
    _, data = run(client, "image-converter", [("x.png", make_image("PNG"), "image/png")], {"format": "ppm"})
    assert Image.open(io.BytesIO(data)).format == "PPM"


def test_image_converter_multi_and_resize(client):
    files = [("a.png", make_image("PNG"), "image/png"), ("b.png", make_image("PNG"), "image/png")]
    _, data = run(client, "image-converter", files, {"format": "gif"})
    assert len(zipfile.ZipFile(io.BytesIO(data)).namelist()) == 2
    _, data = run(client, "resize-image", [("a.png", make_image("PNG"), "image/png")], {"width": "160"})
    assert Image.open(io.BytesIO(data)).size == (160, 100)
    record, _ = run(client, "compress-image", [("a.jpg", make_image("JPEG", mode="RGB"), "image/jpeg")])
    assert record["meta"]["compressed_size"] <= record["meta"]["original_size"]


@pytest.mark.skipif(not engines.is_available(engines.LIBREOFFICE), reason="LibreOffice not installed")
def test_word_to_pdf(client):
    doc = DocxDocument()
    doc.add_paragraph("Converted by LibreOffice")
    buf = io.BytesIO()
    doc.save(buf)
    _, data = run(client, "word-to-pdf", [("letter.docx", buf.getvalue(), "application/octet-stream")])
    assert "Converted by LibreOffice" in open_bytes(data)[0].get_text()


# ---------------------------------------------------------------- edit


def test_rotate(client):
    _, data = run(client, "rotate-pdf", [pdf_file(pages=3)], {"angle": "90", "pages": "2"})
    doc = open_bytes(data)
    assert [p.rotation for p in doc] == [0, 90, 0]


def test_page_numbers(client):
    _, data = run(client, "add-page-numbers", [pdf_file(pages=3)], {"format": "page-n-of-total"})
    assert "Page 2 of 3" in open_bytes(data)[1].get_text()


def test_watermark_text_and_image(client):
    _, data = run(client, "add-watermark", [pdf_file(pages=1)], {"text": "CONFIDENTIAL"})
    assert "CONFIDENTIAL" in open_bytes(data)[0].get_text()
    _, data = run(client, "add-watermark", [pdf_file(pages=1)], {"position": "bottom-right"},
                  extras={"image": ("logo.png", make_image("PNG"), "image/png")})
    assert open_bytes(data)[0].get_images()


def _visual_direction(page, needle):
    line = next(l for b in page.get_text("dict")["blocks"] for l in b.get("lines", [])
                if needle in "".join(s["text"] for s in l["spans"]))
    m = page.rotation_matrix  # unrotated -> what the reader sees
    return pymupdf.Point(line["dir"]) * pymupdf.Matrix(m.a, m.b, m.c, m.d, 0, 0)


@pytest.mark.parametrize("page_rotation", [0, 90, 180, 270])
def test_watermark_orientation_follows_reader(client, page_rotation):
    src = pymupdf.open(stream=make_pdf(pages=1), filetype="pdf")
    src[0].set_rotation(page_rotation)
    _, data = run(client, "add-watermark", [("r.pdf", src.tobytes(), PDF)], {"text": "FLAT", "rotation": "0"})
    assert _visual_direction(open_bytes(data)[0], "FLAT").x > 0.99
    _, data = run(client, "add-watermark", [("r.pdf", src.tobytes(), PDF)], {"text": "TILT", "rotation": "45"})
    d = _visual_direction(open_bytes(data)[0], "TILT")
    assert d.x > 0.6 and d.y < -0.6  # rising to the right, as the reader sees it


def test_page_numbers_on_rotated_page(client):
    src = pymupdf.open(stream=make_pdf(pages=1), filetype="pdf")
    src[0].set_rotation(90)
    _, data = run(client, "add-page-numbers", [("r.pdf", src.tobytes(), PDF)], {"custom_format": "Sheet {n}"})
    page = open_bytes(data)[0]
    assert _visual_direction(page, "Sheet 1").x > 0.99
    hit = page.search_for("Sheet 1")[0] * page.rotation_matrix
    assert hit.y0 > page.rect.height * 0.8  # bottom of the page as the reader sees it


def test_crop(client):
    _, data = run(client, "crop-pdf", [pdf_file(pages=1)], {"top": "10", "bottom": "10", "left": "5", "right": "5"})
    page = open_bytes(data)[0]
    assert round(page.rect.width) == round(595 * 0.9) and round(page.rect.height) == round(842 * 0.8)


def test_edit_pdf(client):
    ops = [
        {"type": "text", "page": 1, "x": 72, "y": 200, "text": "Added note", "size": 14, "color": "#cc0000"},
        {"type": "rect", "page": 1, "x": 60, "y": 190, "width": 200, "height": 40, "stroke": "#000000"},
        {"type": "highlight", "page": 1, "x": 72, "y": 80, "width": 150, "height": 25},
        {"type": "ink", "page": 2, "points": [[10, 10], [50, 60], [90, 20]]},
        {"type": "image", "page": 1, "x": 300, "y": 500, "width": 100, "height": 60, "image": 0},
        {"type": "link", "page": 1, "x": 72, "y": 600, "width": 100, "height": 20, "url": "https://example.com"},
    ]
    _, data = run(client, "edit-pdf", [pdf_file(pages=2)], {"operations": json.dumps(ops)},
                  extras={"images": ("pic.png", make_image("PNG"), "image/png")})
    page = open_bytes(data)[0]
    assert "Added note" in page.get_text()
    assert page.get_links()[0]["uri"] == "https://example.com"


def test_forms(client):
    inspected = client.post("/api/v1/inspect", files={"file": ("form.pdf", make_form_pdf(), PDF)}).json()
    assert {f["name"] for f in inspected["form_fields"]} == {"full_name", "agree"}
    fields = {"full_name": "Talha Khan", "agree": True}
    _, data = run(client, "pdf-forms", [("form.pdf", make_form_pdf(), PDF)], {"fields": json.dumps(fields)})
    widgets = {w.field_name: w.field_value for w in open_bytes(data)[0].widgets()}
    assert widgets["full_name"] == "Talha Khan" and widgets["agree"] not in ("Off", False, "")


# ---------------------------------------------------------------- security


def test_protect_then_unlock(client):
    _, locked = run(client, "protect-pdf", [pdf_file(pages=1)], {"new_password": "s3cret!"})
    assert open_bytes(locked).needs_pass
    err = run(client, "unlock-pdf", [("locked.pdf", locked, PDF)], expect=422)
    assert err["error"]["code"] == "password_required"
    err = run(client, "unlock-pdf", [("locked.pdf", locked, PDF)], {"password": "nope"}, expect=422)
    assert err["error"]["code"] == "wrong_password"
    _, unlocked = run(client, "unlock-pdf", [("locked.pdf", locked, PDF)], {"password": "s3cret!"})
    assert not open_bytes(unlocked).needs_pass


def test_sign(client):
    placements = [{"page": 1, "x": 350, "y": 700, "width": 180, "height": 60},
                  {"page": 1, "x": 350, "y": 770, "width": 180, "height": 20, "kind": "date"}]
    _, data = run(client, "sign-pdf", [pdf_file(pages=1)], {"placements": json.dumps(placements)},
                  extras={"signature": ("sig.png", make_image("PNG"), "image/png")})
    assert open_bytes(data)[0].get_images()
    _, data = run(client, "sign-pdf", [pdf_file(pages=1)],
                  {"placements": json.dumps(placements[:1]), "typed_name": "Talha Khan"})
    assert "Talha Khan" in open_bytes(data)[0].get_text()


def test_redact(client):
    record, data = run(client, "redact-pdf", [pdf_file(pages=2)],
                       {"terms": json.dumps(["Hello page 1"]), "patterns": "email"})
    text = "".join(p.get_text() for p in open_bytes(data))
    assert "Hello page 1" not in text and "someone@example.com" not in text
    assert "Hello page 2" in text
    assert record["meta"]["redactions"] >= 3


def test_compare(client):
    record, data = run(client, "compare-pdf",
                       [pdf_file("old.pdf", pages=2), pdf_file("new.pdf", pages=2, text="Changed page")])
    assert open_bytes(data).page_count == 2
    assert record["meta"]["words_added"] > 0 and not record["meta"]["identical"]


# ---------------------------------------------------------------- engine-dependent / AI


def test_ai_tools_report_unavailable_without_key(client):
    tool = client.get("/api/v1/tools/ai-summarizer").json()
    assert tool["available"] is False
    resp = client.post("/api/v1/convert/ai-summarizer", files=[("files", pdf_file())])
    assert resp.status_code == 503


@pytest.fixture
def fake_claude(monkeypatch):
    """Stand-in for the Claude API so the PDF side of the AI tools can be tested offline."""
    from app.tools import ai

    calls = []

    def fake(system, prompt, *, effort, schema=None, max_tokens=16000):
        calls.append({"prompt": prompt, "schema": schema})
        if schema:  # translation request
            items = json.loads(prompt)["items"]
            return json.dumps({"items": [{"id": i["id"], "text": "ES " + i["text"]} for i in items]})
        return "## Summary\n- It says hello."

    monkeypatch.setattr(engines, "ai_available", lambda: True)
    monkeypatch.setattr(ai, "ask_claude", fake)
    return calls


def test_ai_summarizer(client, fake_claude):
    record, data = run(client, "ai-summarizer", [pdf_file(pages=2)], {"length": "short"})
    assert data.decode().startswith("## Summary")
    assert record["meta"]["summary"].startswith("## Summary")
    assert "Hello page 2" in fake_claude[0]["prompt"]


def test_translate_keeps_layout(client, fake_claude):
    _, data = run(client, "translate-pdf", [pdf_file(pages=2)], {"target_language": "Spanish"})
    doc = open_bytes(data)
    text = doc[0].get_text()
    assert "ES Hello page 1" in text
    assert "Hello page 1" not in text.replace("ES Hello page 1", "")  # original text was replaced
    assert doc[0].get_drawings()  # graphics (the blue box) were kept


# ---------------------------------------------------------------- validation & lifecycle


def test_rejects_wrong_type_and_fake_pdf(client):
    err = run(client, "pdf-to-word", [("x.png", make_image("PNG"), "image/png")], expect=422)
    assert err["error"]["code"] == "unsupported_file_type"
    err = run(client, "pdf-to-word", [("fake.pdf", b"not a pdf at all", PDF)], expect=422)
    assert err["error"]["code"] == "invalid_file"
    err = run(client, "merge-pdf", [pdf_file()], expect=422)
    assert err["error"]["code"] == "no_files"


def test_bad_options(client):
    err = run(client, "pdf-to-jpg", [pdf_file()], {"dpi": "5000"}, expect=422)
    assert "between" in err["error"]["message"]
    err = run(client, "split-pdf", [pdf_file(pages=2)], {"ranges": "1-9"}, expect=422)
    assert "outside" in err["error"]["message"]


def test_unknown_tool_and_missing_job(client):
    assert client.post("/api/v1/convert/nope", files=[("files", pdf_file())]).status_code == 404
    assert client.get("/api/v1/download/doesnotexistdoesnotexist").status_code == 404
    assert client.get("/api/v1/download/../../etc").status_code == 404


def test_delete_job(client):
    record, _ = run(client, "rotate-pdf", [pdf_file()])
    assert client.delete(f"/api/v1/jobs/{record['job_id']}").status_code == 200
    assert client.get(record["download_url"]).status_code == 404


def test_rate_limit():
    from app.core.errors import RateLimited
    from app.core.rate_limit import RateLimiter

    limiter = RateLimiter(2)
    limiter.check("1.2.3.4")
    limiter.check("1.2.3.4")
    with pytest.raises(RateLimited):
        limiter.check("1.2.3.4")
    limiter.check("5.6.7.8")
